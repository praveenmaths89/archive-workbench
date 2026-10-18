"""Exporters: DOCX, searchable PDF (original image + invisible text), TXT, MD, JSON."""
from __future__ import annotations

import io

import pymupdf
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

from .ir import DocumentIR


def done_pages(ir: DocumentIR):
    """Pages that were actually OCR'd (or failed); skipped pages have nothing to export."""
    return [p for p in ir.pages if p.status in ("done", "failed")]


def to_txt(ir: DocumentIR) -> str:
    return "\n\n".join(f"--- Page {p.number} ---\n{p.text}" for p in done_pages(ir))


def to_md(ir: DocumentIR) -> str:
    out = [f"# {ir.filename}\n"]
    for p in done_pages(ir):
        out.append(f"\n---\n*Page {p.number}*\n")
        for b in p.blocks:
            if b.type == "table" and b.table_rows():
                rows = b.table_rows()
                width = max(len(r) for r in rows)
                rows = [r + [""] * (width - len(r)) for r in rows]
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("|" + "---|" * width)
                out.extend("| " + " | ".join(c.replace("|", "/") for c in r) + " |" for r in rows[1:])
                out.append("")
                continue
            out.append(f"## {b.text}" if b.type == "heading" else
                       f"- {b.text}" if b.type == "list_item" else
                       f"> {b.text}" if b.type == "marginalia" else b.text)
            out.append("")
    return "\n".join(out)


def to_json(ir: DocumentIR) -> str:
    return ir.model_dump_json(indent=2)


def to_docx(ir: DocumentIR, page_breaks: bool = True) -> bytes:
    d = Document()
    d.add_heading(ir.filename, level=0)
    pages = done_pages(ir)
    for i, p in enumerate(pages):
        if p.status == "failed":
            d.add_paragraph(f"[Page {p.number} could not be processed: {p.error}]")
        for b in p.blocks:
            if b.type == "heading":
                d.add_heading(b.text, level=2)
            elif b.type == "list_item":
                d.add_paragraph(b.text, style="List Bullet")
            elif b.type == "table" and b.table_rows():
                rows = b.table_rows()
                width = max(len(r) for r in rows)
                t = d.add_table(rows=len(rows), cols=width)
                t.style = "Table Grid"
                for ri, r in enumerate(rows):
                    for ci in range(width):
                        t.cell(ri, ci).text = r[ci] if ci < len(r) else ""
                d.add_paragraph()
            elif b.type == "figure":
                para = d.add_paragraph()
                run = para.add_run(f"[Figure] {b.text}")
                run.italic = True
            elif b.type == "marginalia":
                para = d.add_paragraph()
                run = para.add_run(f"[Margin: {b.text}]")
                run.italic = True
                run.font.size = Pt(9)
            else:
                d.add_paragraph(b.text)
        if page_breaks and i < len(pages) - 1:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_searchable_pdf(ir: DocumentIR) -> bytes:
    """Original page image with the transcription as an invisible text layer.

    Limitation (M1): the model returns no coordinates, so the text layer is laid out over the
    full page rather than aligned to words. Search finds the page; highlight is approximate.
    """
    out = pymupdf.open()
    for p in done_pages(ir):
        if not p.image_path:
            continue
        img = pymupdf.open(p.image_path)
        rect = img[0].rect
        page = out.new_page(width=rect.width, height=rect.height)
        page.insert_image(page.rect, filename=p.image_path)
        if p.text.strip():
            page.insert_textbox(page.rect + (10, 10, -10, -10), p.text, fontsize=6,
                                render_mode=3, fontname="helv")
        img.close()
    data = out.tobytes(garbage=3, deflate=True)
    out.close()
    return data


def markdown_to_docx(md: str, title: str) -> bytes:
    """Minimal renderer for research reports: headings, bullets, paragraphs, **bold**."""
    import re
    d = Document()
    d.add_heading(title, level=0)
    for line in md.splitlines():
        line = line.rstrip()
        if not line:
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            d.add_heading(m.group(2), level=min(len(m.group(1)), 4))
            continue
        bullet = re.match(r"^\s*[-*]\s+(.*)", line)
        para = d.add_paragraph(style="List Bullet") if bullet else d.add_paragraph()
        text = bullet.group(1) if bullet else line
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if part:
                para.add_run(part).bold = bool(i % 2)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
