import json
import time
from types import SimpleNamespace as NS

import pymupdf
import pytest
from fastapi.testclient import TestClient

from app import export, main, research
from app.adapters import result_from_json
from app.config import settings
from app.ir import Block, DocumentIR, Page
from app.pipeline import count_uncertain, needs_review
from tests.test_pipeline import make_scan_pdf, wait_done


@pytest.fixture(autouse=True)
def tmp_data(tmp_path, monkeypatch):
    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr(settings, "use_mock", True)


def test_table_rows_parsed_from_model_json_and_pipe_fallback():
    r = result_from_json({"blocks": [{"type": "table", "text": "", "rows": [["Date", "Init"], ["June 22", ""]]}]}, 1, 1)
    assert r.blocks[0].table_rows() == [["Date", "Init"], ["June 22", ""]]
    b = Block(type="table", text="| A | B |\n|---|---|\n| 1 | 2 |")
    assert b.table_rows() == [["A", "B"], ["1", "2"]]


def test_docx_has_real_table_and_md_pipe_table():
    ir = DocumentIR(doc_id="d", filename="x.pdf", pages=[Page(number=1, status="done", blocks=[
        Block(type="table", rows=[["a", "b"], ["c", "d"]]), Block(type="figure", text="A chart")])])
    import io
    from docx import Document
    d = Document(io.BytesIO(export.to_docx(ir)))
    assert len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "d"
    assert "| a | b |" in export.to_md(ir) and "|---|---|" in export.to_md(ir)


def test_review_flag_from_uncertain_markers_not_self_reported_confidence():
    good = Page(number=1, confidence=0.78, blocks=[Block(text="Clear text with one [word?] only")])
    good.uncertain = count_uncertain(good)
    assert good.uncertain == 1 and not needs_review(good)  # 0.78 alone no longer flags
    bad = Page(number=2, confidence=0.9, blocks=[Block(text="[illegible] a [x?] b [illegible]")])
    bad.uncertain = count_uncertain(bad)
    assert bad.uncertain == 3 and needs_review(bad)
    assert needs_review(Page(number=3, confidence=0.4))


def test_page_image_edit_and_research_guard():
    c = TestClient(main.app)
    r = c.post("/api/v1/jobs", files=[("files", ("s.pdf", make_scan_pdf(1), "application/pdf"))])
    job = r.json()["job_id"]
    doc = wait_done(c, job)[0]["id"]
    img = c.get(f"/api/v1/jobs/{job}/documents/{doc}/pages/1/image")
    assert img.status_code == 200 and img.content[:2] == b"\xff\xd8"
    assert c.get(f"/api/v1/jobs/{job}/documents/{doc}/pages/9/image").status_code == 404

    e = c.put(f"/api/v1/jobs/{job}/documents/{doc}/pages/1",
              json={"blocks": [{"type": "paragraph", "text": "Corrected [word?]"},
                               {"type": "table", "text": "", "rows": [["x", "y"]]}]})
    assert e.status_code == 200
    page = c.get(f"/api/v1/jobs/{job}/documents/{doc}").json()["pages"][0]
    assert page["edited"] and page["blocks"][1]["rows"] == [["x", "y"]] and page["uncertain"] == 1
    ir = main.load_ir(job, doc)
    assert ir.pages[0].original_blocks and "mock transcription" in ir.pages[0].original_blocks[0].text
    # research refuses in mock mode instead of silently doing nothing
    assert c.post(f"/api/v1/jobs/{job}/documents/{doc}/research", json={}).status_code == 400


class FakeStream:
    def __init__(self, msg): self.msg = msg
    def __enter__(self): return self
    def __exit__(self, *a): pass
    def get_final_message(self): return self.msg


class FakeClient:
    def __init__(self, msgs):
        self.msgs, self.calls = list(msgs), []
        self.messages = NS(stream=self._stream)
    def _stream(self, **kw):
        self.calls.append(kw)
        return FakeStream(self.msgs.pop(0))


def usage(): return NS(input_tokens=100, output_tokens=50)


def test_research_cites_sources_records_queries_and_resumes_pause_turn():
    hit = NS(url="https://archive.example/a", title="Archive A")
    first = NS(stop_reason="pause_turn", usage=usage(), content=[
        NS(type="server_tool_use", input={"query": "Wingate chaplain Belgaum 1885"}, model_dump=lambda **k: {"type": "server_tool_use"}),
        NS(type="server_tool_use", input={"code": "print(1)"}, model_dump=lambda **k: {"type": "server_tool_use"}),
        NS(type="web_search_tool_result", content=[hit, NS(url="https://other.example/b", title="Other B")],
           model_dump=lambda **k: {"type": "web_search_tool_result"})])
    second = NS(stop_reason="end_turn", usage=usage(), content=[
        NS(type="text", text="He was appointed in 1885.", citations=[NS(url="https://archive.example/a", title="Archive A")],
           model_dump=lambda **k: {"type": "text"})])
    fc = FakeClient([first, second])
    res = research.run_research("Minute", "doc text", "Who?", "quick", allowed_domains=["archive.example"], client=fc)
    assert res.answer_md == "He was appointed in 1885. [1]"
    assert [s.url for s in res.sources] == ["https://archive.example/a"]
    assert [s.url for s in res.consulted] == ["https://other.example/b"]  # seen, not cited
    assert res.queries == ["Wingate chaplain Belgaum 1885"]  # code blocks are not queries
    assert len(fc.calls) == 2 and fc.calls[0]["tools"][0]["allowed_domains"] == ["archive.example"]
    assert fc.calls[0]["tools"][0]["type"] == "web_search_20250305" and fc.calls[0]["tools"][0]["max_uses"] == 4
    assert fc.calls[1]["messages"][-1]["role"] == "assistant"  # resumed without a fake 'Continue.' turn
    assert res.input_tokens == 200


def test_research_no_results_is_reported_not_invented():
    msg = NS(stop_reason="end_turn", usage=usage(), content=[
        NS(type="text", text="I could not find reliable sources.", citations=None, model_dump=lambda **k: {})])
    res = research.run_research("D", "t", None, client=FakeClient([msg]))
    assert "could not find" in res.answer_md and res.sources == []


def test_research_rejects_both_domain_lists_and_bad_depth():
    with pytest.raises(ValueError):
        research.run_research("D", "t", allowed_domains=["a.com"], blocked_domains=["b.com"], client=FakeClient([]))
    with pytest.raises(ValueError):
        research.run_research("D", "t", depth="nope", client=FakeClient([]))


def test_research_endpoint_flow_and_download(monkeypatch):
    monkeypatch.setattr(settings, "use_mock", False)
    monkeypatch.setattr(settings, "anthropic_api_key", "test-key")
    monkeypatch.setattr(research, "run_research", lambda *a, **k: research.ResearchResult(
        answer_md="## Summary\nFacts [1]", sources=[research.Source(1, "https://x.example", "X")], queries=["q1"]))
    monkeypatch.setattr(main, "get_adapter", lambda: __import__("app.adapters", fromlist=["x"]).MockAdapter())
    c = TestClient(main.app)
    job = c.post("/api/v1/jobs", files=[("files", ("s.pdf", make_scan_pdf(1), "application/pdf"))]).json()["job_id"]
    doc = wait_done(c, job)[0]["id"]
    rid = c.post(f"/api/v1/jobs/{job}/documents/{doc}/research", json={"question": "Q?"}).json()["research_id"]
    for _ in range(50):
        rec = c.get(f"/api/v1/research/{rid}").json()
        if rec["status"] != "running":
            break
        time.sleep(0.1)
    assert rec["status"] == "done" and rec["sources"][0]["n"] == 1
    md = c.get(f"/api/v1/research/{rid}/download?format=md").text
    assert "Facts [1]" in md and "[X](https://x.example)" in md and "- q1" in md
    assert c.get(f"/api/v1/research/{rid}/download?format=docx").content[:2] == b"PK"
    assert c.get(f"/api/v1/research/nope/download").status_code == 404


def test_inline_markdown_links_become_sources():
    msg = NS(stop_reason="end_turn", usage=usage(), content=[
        NS(type="text", text="See [Hansard](https://hansard.example/x) for details.", citations=None, model_dump=lambda **k: {})])
    res = research.run_research("D", "t", None, client=FakeClient([msg]))
    assert [s.url for s in res.sources] == ["https://hansard.example/x"]


# ---------- page selection ----------
from app.ranges import parse_pages


def test_parse_pages_forms_and_errors():
    assert parse_pages("1-3, 5, 8-", 10) == [1, 2, 3, 5, 8, 9, 10]
    assert parse_pages("-3", 10) == [1, 2, 3]
    assert parse_pages("7-99", 10) == [7, 8, 9, 10]  # end past the last page is clamped
    assert parse_pages("", 10) is None and parse_pages("all", 10) is None
    for bad in ("0", "11", "5-2", "abc", "1-3,x"):
        with pytest.raises(ValueError):
            parse_pages(bad, 10)


def _upload(c, pages_spec=None, by_file=None, n=6, name="scan.pdf"):
    opts = {}
    if pages_spec is not None:
        opts["pages"] = pages_spec
    if by_file:
        opts["pages_by_file"] = by_file
    r = c.post("/api/v1/jobs", files=[("files", (name, make_scan_pdf(n), "application/pdf"))],
               data={"options": json.dumps(opts)})
    assert r.status_code == 200
    return r.json()["job_id"], r.json()["documents"][0]


def test_only_selected_pages_are_ocrd_and_others_stay_skipped(monkeypatch):
    calls = []
    from app.adapters import MockAdapter
    ad = MockAdapter()
    monkeypatch.setattr(main, "get_adapter", lambda: ad)
    c = TestClient(main.app)
    job, d = _upload(c, "2-3, 5")
    assert d["pages"] == 3
    doc = wait_done(c, job)[0]
    assert doc["status"] == "done" and doc["total_pages"] == 6 and doc["selected_pages"] == 3
    assert doc["done_pages"] == 3 and len(ad.calls) == 3  # exactly 3 model calls, not 6
    pages = c.get(f"/api/v1/jobs/{job}/documents/{doc['id']}").json()["pages"]
    assert [p["status"] for p in pages] == ["skipped", "done", "done", "skipped", "done", "skipped"]
    # exports contain only OCR'd pages
    txt = c.get(f"/api/v1/jobs/{job}/documents/{doc['id']}/download?format=txt").text
    assert "Page 2" in txt and "Page 1 ---" not in txt and "Page 6" not in txt
    pdf = pymupdf.open(stream=c.get(f"/api/v1/jobs/{job}/documents/{doc['id']}/download?format=pdf").content)
    assert pdf.page_count == 3
    # a skipped page can still be previewed
    assert c.get(f"/api/v1/jobs/{job}/documents/{doc['id']}/pages/6/image").content[:2] == b"\xff\xd8"


def test_ocr_more_pages_later_keeps_done_pages(monkeypatch):
    from app.adapters import MockAdapter
    ad = MockAdapter()
    monkeypatch.setattr(main, "get_adapter", lambda: ad)
    c = TestClient(main.app)
    job, _ = _upload(c, "1")
    doc = wait_done(c, job)[0]
    assert len(ad.calls) == 1
    r = c.post(f"/api/v1/jobs/{job}/documents/{doc['id']}/ocr", json={"pages": "1-3"})
    assert r.json() == {"status": "queued", "pages": 2}  # page 1 already done, so only 2 new
    time.sleep(0.2)
    doc = wait_done(c, job)[0]
    assert len(ad.calls) == 3 and doc["done_pages"] == 3
    # nothing to do when everything requested is already done
    assert c.post(f"/api/v1/jobs/{job}/documents/{doc['id']}/ocr", json={"pages": "1-3"}).json()["status"] == "nothing_to_do"
    # redo re-runs a done page
    c.post(f"/api/v1/jobs/{job}/documents/{doc['id']}/ocr", json={"pages": "2", "redo": True})
    time.sleep(0.2)
    wait_done(c, job)
    assert len(ad.calls) == 4
    assert c.post(f"/api/v1/jobs/{job}/documents/{doc['id']}/ocr", json={"pages": "99"}).status_code == 400


def test_per_file_range_and_bad_range_rejects_only_that_file(monkeypatch):
    from app.adapters import MockAdapter
    monkeypatch.setattr(main, "get_adapter", lambda: MockAdapter())
    c = TestClient(main.app)
    files = [("files", ("a.pdf", make_scan_pdf(4), "application/pdf")),
             ("files", ("b.pdf", make_scan_pdf(4), "application/pdf")),
             ("files", ("c.pdf", make_scan_pdf(4), "application/pdf"))]
    r = c.post("/api/v1/jobs", files=files, data={"options": json.dumps(
        {"pages": "1", "pages_by_file": {"b.pdf": "2-4", "c.pdf": "9"}})})
    docs = {d["filename"]: d for d in r.json()["documents"]}
    assert docs["a.pdf"]["pages"] == 1 and docs["b.pdf"]["pages"] == 3
    assert docs["c.pdf"]["status"] == "rejected" and "page 9 is outside 1-4" in docs["c.pdf"]["error"]


# ---------- start / stop ----------
def test_stop_mid_document_keeps_finished_pages_and_resume_finishes_rest(monkeypatch):
    from app.adapters import MockAdapter

    class Slow(MockAdapter):
        def ocr(self, *a, **k):
            time.sleep(0.4)
            return super().ocr(*a, **k)

    ad = Slow()
    monkeypatch.setattr(main, "get_adapter", lambda: ad)
    c = TestClient(main.app)
    job, d = _upload(c, None, n=6)
    did = d["doc_id"]
    while len(ad.calls) < 2:
        time.sleep(0.05)
    assert c.post(f"/api/v1/jobs/{job}/documents/{did}/stop").json()["status"] == "stopping"
    doc = wait_done(c, job)[0]
    assert doc["status"] == "stopped"
    done_at_stop = len(ad.calls)
    assert 2 <= done_at_stop < 6  # stopped before finishing, after the page in flight
    time.sleep(0.8)
    assert len(ad.calls) == done_at_stop  # nothing keeps running in the background
    pages = c.get(f"/api/v1/jobs/{job}/documents/{did}").json()["pages"]
    statuses = [p["status"] for p in pages]
    assert statuses.count("done") == done_at_stop and "pending" not in statuses
    # resume = start OCR on everything not done yet
    r = c.post(f"/api/v1/jobs/{job}/documents/{did}/ocr", json={})
    assert r.json()["pages"] == 6 - done_at_stop
    time.sleep(0.2)
    assert wait_done(c, job, timeout=30)[0]["status"] == "done" and len(ad.calls) == 6
    assert c.post(f"/api/v1/jobs/{job}/documents/{did}/stop").json()["note"] == "not running"


def test_cannot_start_ocr_while_running(monkeypatch):
    from app.adapters import MockAdapter

    class Slow(MockAdapter):
        def ocr(self, *a, **k):
            time.sleep(0.3)
            return super().ocr(*a, **k)

    monkeypatch.setattr(main, "get_adapter", lambda: Slow())
    c = TestClient(main.app)
    job, d = _upload(c, None, n=3)
    assert c.post(f"/api/v1/jobs/{job}/documents/{d['doc_id']}/ocr", json={}).status_code == 409
    c.post(f"/api/v1/jobs/{job}/documents/{d['doc_id']}/stop")
    wait_done(c, job)


def test_billing_error_stops_the_whole_run_with_a_clear_message(monkeypatch):
    from app.adapters import MockAdapter

    class NoCredit(MockAdapter):
        def ocr(self, *a, **k):
            self.calls.append("x")
            raise RuntimeError("Error code: 400 - Your credit balance is too low to access the Anthropic API.")

    ad = NoCredit()
    monkeypatch.setattr(main, "get_adapter", lambda: ad)
    c = TestClient(main.app)
    job, d = _upload(c, None, n=5)
    doc = wait_done(c, job)[0]
    assert doc["status"] == "failed" and "credit balance is too low" in doc["error"].lower()
    assert "Add credits" in doc["error"]
    assert len(ad.calls) == 1  # did not hammer the API for every page
    pages = c.get(f"/api/v1/jobs/{job}/documents/{d['doc_id']}").json()["pages"]
    assert {p["status"] for p in pages} == {"skipped"}  # nothing stuck as 'pending' or 'failed'
