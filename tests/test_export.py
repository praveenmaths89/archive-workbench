import io
import xml.etree.ElementTree as ET
from datetime import datetime, timezone

from docx import Document

from archive_workbench.export import to_bibtex, to_docx, to_md, to_tei
from archive_workbench.ir import Block, DocumentIR, Page


def _ir() -> DocumentIR:
    return DocumentIR(doc_id="d", filename="minute.pdf", pages=[
        Page(number=1, status="done", blocks=[
            Block(type="paragraph", text="Kittur revenue"),
            Block(type="table", rows=[["Year", "Rupees"], ["1862", "400"]]),
        ]),
    ])


def test_docx_has_a_real_table_and_markdown_has_pipes():
    blob = to_docx(_ir())
    assert blob[:2] == b"PK"
    doc = Document(io.BytesIO(blob))
    assert doc.tables[0].cell(1, 1).text == "400"
    md = to_md(_ir())
    assert "| Year | Rupees |" in md
    assert "Kittur revenue" in md


def test_bibtex_and_tei():
    bib = to_bibtex("minute.pdf", "Minute Paper", datetime(2026, 1, 2, tzinfo=timezone.utc))
    assert bib.startswith("@unpublished{")
    assert "Minute Paper" in bib
    xml = to_tei("minute.pdf", "Minute Paper", [{
        "number": 1,
        "translation": "[mock translation to hi]\nKittur",
        "translation_language": "hi",
        "regions": [{"type": "line", "text": "Kittur revenue", "x": 1, "y": 2, "w": 3, "h": 4}],
    }])
    root = ET.fromstring(xml)
    assert root.tag.endswith("TEI")
    texts = [el.text for el in root.iter() if el.text]
    assert "Kittur revenue" in texts
    facs = [el.attrib.get("facs") for el in root.iter() if "facs" in el.attrib]
    assert "1,2,3,4" in facs
