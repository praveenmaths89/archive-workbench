"""Exporters. DOCX, Markdown, JSON, TXT, and searchable PDF come from the v0.2 prototype.
BibTeX and TEI are new.
"""
from __future__ import annotations

import io
import re
import xml.etree.ElementTree as ET
from datetime import datetime

import pymupdf
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt

from .ir import Block, DocumentIR

TEI_NS = "http://www.tei-c.org/ns/1.0"


def done_pages(ir: DocumentIR):
    return [p for p in ir.pages if p.status in ("done", "failed")]


def to_txt(ir: DocumentIR) -> str:
    return "\n\n".join(f"--- Page {p.number} ---\n{p.text}" for p in done_pages(ir))


def to_md(ir: DocumentIR) -> str:
    out = [f"# {ir.filename}\n"]
    for p in done_pages(ir):
        out.append(f"\n---\n*Page {p.number}*\n")
        for b in p.blocks:
            if b.type == "table" and b.table_rows():
                rows = b.table_rows()
                width = max(len(r) for r in rows)
                rows = [r + [""] * (width - len(r)) for r in rows]
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("|" + "---|" * width)
                out.extend("| " + " | ".join(c.replace("|", "/") for c in r) + " |" for r in rows[1:])
                out.append("")
                continue
            out.append(f"## {b.text}" if b.type == "heading" else
                       f"- {b.text}" if b.type == "list_item" else
                       f"> {b.text}" if b.type == "marginalia" else b.text)
            out.append("")
    return "\n".join(out)


def to_json(ir: DocumentIR) -> str:
    return ir.model_dump_json(indent=2)


def to_docx(ir: DocumentIR, page_breaks: bool = True) -> bytes:
    d = Document()
    d.add_heading(ir.filename, level=0)
    pages = done_pages(ir)
    for i, p in enumerate(pages):
        if p.status == "failed":
            d.add_paragraph(f"[Page {p.number} could not be processed: {p.error}]")
        for b in p.blocks:
            if b.type == "heading":
                d.add_heading(b.text, level=2)
            elif b.type == "list_item":
                d.add_paragraph(b.text, style="List Bullet")
            elif b.type == "table" and b.table_rows():
                rows = b.table_rows()
                width = max(len(r) for r in rows)
                t = d.add_table(rows=len(rows), cols=width)
                t.style = "Table Grid"
                for ri, r in enumerate(rows):
                    for ci in range(width):
                        t.cell(ri, ci).text = r[ci] if ci < len(r) else ""
                d.add_paragraph()
            elif b.type == "figure":
                para = d.add_paragraph()
                run = para.add_run(f"[Figure] {b.text}")
                run.italic = True
            elif b.type == "marginalia":
                para = d.add_paragraph()
                run = para.add_run(f"[Margin: {b.text}]")
                run.italic = True
                run.font.size = Pt(9)
            else:
                d.add_paragraph(b.text)
        if page_breaks and i < len(pages) - 1:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def to_searchable_pdf(ir: DocumentIR) -> bytes:
    """Page image plus an invisible text layer. The layer covers the page; it is not word-aligned."""
    out = pymupdf.open()
    for p in done_pages(ir):
        if not p.image_path:
            continue
        img = pymupdf.open(p.image_path)
        rect = img[0].rect
        page = out.new_page(width=rect.width, height=rect.height)
        page.insert_image(page.rect, filename=p.image_path)
        if p.text.strip():
            page.insert_textbox(page.rect + (10, 10, -10, -10), p.text, fontsize=6,
                                render_mode=3, fontname="helv")
        img.close()
    data = out.tobytes(garbage=3, deflate=True)
    out.close()
    return data


def markdown_to_docx(md: str, title: str) -> bytes:
    d = Document()
    d.add_heading(title, level=0)
    for line in md.splitlines():
        line = line.rstrip()
        if not line:
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            d.add_heading(m.group(2), level=min(len(m.group(1)), 4))
            continue
        bullet = re.match(r"^\s*[-*]\s+(.*)", line)
        para = d.add_paragraph(style="List Bullet") if bullet else d.add_paragraph()
        text = bullet.group(1) if bullet else line
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if part:
                para.add_run(part).bold = bool(i % 2)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _bib_escape(value: str) -> str:
    return value.replace("\\", "").replace("{", "(").replace("}", ")").replace("\n", " ").strip()


def to_bibtex(filename: str, title: str, created: datetime | None) -> str:
    year = created.year if created else datetime.now().year
    key = re.sub(r"[^A-Za-z0-9]+", "", PathStem(filename))[:40] or "document"
    return (
        f"@unpublished{{{key}{year},\n"
        f"  title = {{{_bib_escape(title or filename)}}},\n"
        f"  year = {{{year}}},\n"
        f"  note = {{Transcribed locally with Archive Workbench.}}\n"
        f"}}\n"
    )


def PathStem(filename: str) -> str:
    return filename.rsplit(".", 1)[0] if "." in filename else filename


def to_tei(filename: str, title: str, pages: list[dict]) -> bytes:
    """TEI P5. `pages` items have number, regions (text, type, x, y, w, h), and optional translation."""
    ET.register_namespace("", TEI_NS)
    root = ET.Element(f"{{{TEI_NS}}}TEI")
    header = ET.SubElement(root, f"{{{TEI_NS}}}teiHeader")
    file_desc = ET.SubElement(header, f"{{{TEI_NS}}}fileDesc")
    title_stmt = ET.SubElement(file_desc, f"{{{TEI_NS}}}titleStmt")
    ET.SubElement(title_stmt, f"{{{TEI_NS}}}title").text = title or filename
    pub = ET.SubElement(file_desc, f"{{{TEI_NS}}}publicationStmt")
    ET.SubElement(pub, f"{{{TEI_NS}}}p").text = (
        "Export produced locally by Archive Workbench. "
        "This header describes the export file, not the rights in the source document."
    )
    source = ET.SubElement(file_desc, f"{{{TEI_NS}}}sourceDesc")
    ET.SubElement(source, f"{{{TEI_NS}}}p").text = filename
    text = ET.SubElement(root, f"{{{TEI_NS}}}text")
    body = ET.SubElement(text, f"{{{TEI_NS}}}body")
    for page in pages:
        div = ET.SubElement(body, f"{{{TEI_NS}}}div", {"type": "page", "n": str(page["number"])})
        for region in page.get("regions") or []:
            tag = "table" if region.get("type") == "table" else "p"
            el = ET.SubElement(div, f"{{{TEI_NS}}}{tag}")
            box = f"{region.get('x', 0)},{region.get('y', 0)},{region.get('w', 0)},{region.get('h', 0)}"
            el.set("facs", box)
            el.text = region.get("text") or ""
        if page.get("translation"):
            tr = ET.SubElement(div, f"{{{TEI_NS}}}div", {"type": "translation"})
            if page.get("translation_language"):
                tr.set("xml:lang", str(page["translation_language"]))
            ET.SubElement(tr, f"{{{TEI_NS}}}p").text = page["translation"]
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)
